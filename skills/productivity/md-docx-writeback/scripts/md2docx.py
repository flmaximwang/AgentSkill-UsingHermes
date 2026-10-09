#!/usr/bin/env python3
"""把 md 写回既有 docx（沿用母本版式）。见 SKILL.md 的 Procedure。

版式母本自动识别：标题段=最大字号；正文段=其余段中最常见字号；空段=无 run 且无图。
md 解析：'## X'=标题段；'![图注](路径)'=图片段+图注段；其他非空行=正文段；空行仅作分隔。
"""
import argparse
import copy
import datetime as dt
import os
import re
import shutil
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
BLANKS_BETWEEN_SECTIONS = 3  # 与母本一致的分节空段数


def parse_md(md_path):
    blocks = []
    with open(md_path, encoding="utf-8") as fh:
        for raw in fh.read().splitlines():
            s = raw.strip()
            if not s:
                continue
            if s.startswith("## "):
                blocks.append(("heading", s[3:].strip()))
            elif s.startswith("!["):
                m = re.match(r"!\[(.*?)\]\((.*?)\)\s*$", s)
                if not m:
                    raise ValueError("无法解析图片行: %s" % s)
                blocks.append(("image", m.group(2).strip(), m.group(1).strip()))
            else:
                blocks.append(("para", s))
    return blocks


def first_run_rpr(p_el):
    r = p_el.find(W + "r")
    if r is None:
        raise ValueError("模板段落没有 run，无法取字符格式")
    return copy.deepcopy(r.find(W + "rPr"))


def pick_templates(doc):
    """自动挑版式母本段。不能写死段落序号——本脚本的输出会被下一次运行当母本，序号会漂移。"""
    head_p = body_p = blank_p = None
    sizes = {}
    for p in doc.paragraphs:
        has_img = p._p.findall(".//" + W + "drawing")
        if not p.runs and not has_img:
            if blank_p is None:
                blank_p = copy.deepcopy(p._p)
            continue
        if has_img:
            continue
        rpr = p.runs[0]._r.find(W + "rPr")
        sz_el = rpr.find(W + "sz") if rpr is not None else None
        if sz_el is None:
            continue
        sizes.setdefault(int(sz_el.get(W + "val")), []).append(p)

    if not sizes:
        raise ValueError("母本里找不到带字号的段落，无法取版式模板")
    head_sz = max(sizes)
    rest = {sz: ps for sz, ps in sizes.items() if sz != head_sz}
    body_sz = max(rest, key=lambda s: len(rest[s])) if rest else head_sz
    head_p = sizes[head_sz][0]._p
    body_p = sizes[body_sz][0]._p
    if blank_p is None:
        raise ValueError("母本里找不到空段，无法取分节空行模板")
    return head_p, body_p, blank_p, head_sz, body_sz


def build(src, md_path, out_path, img_width_cm, base_dir):
    doc = Document(src)
    body = doc.element.body

    head_src, body_src, blank_p, head_sz, body_sz = pick_templates(doc)
    head_ppr = copy.deepcopy(head_src.find(W + "pPr"))
    head_rpr = first_run_rpr(head_src)
    body_ppr = copy.deepcopy(body_src.find(W + "pPr"))
    body_rpr = first_run_rpr(body_src)
    print("版式母本: 标题段 sz=%s，正文段 sz=%s，空段已定位" % (head_sz, body_sz))

    blocks = parse_md(md_path)
    if not blocks:
        raise ValueError("md 内容为空")

    for p in list(body.findall(W + "p")):
        body.remove(p)
    sect = body.find(W + "sectPr")

    def append_el(el):
        if sect is not None:
            sect.addprevious(el)
        else:
            body.append(el)

    def add_para(ppr, rpr, text=None, align=None):
        p = doc.add_paragraph()  # python-docx 插到 sectPr 之前
        pel = p._p
        old = pel.find(W + "pPr")
        if old is not None:
            pel.remove(old)
        pel.insert(0, copy.deepcopy(ppr))
        if text is not None:
            r = pel.makeelement(W + "r", {})
            r.append(copy.deepcopy(rpr))
            t = pel.makeelement(W + "t", {XML_SPACE: "preserve"})
            t.text = text
            r.append(t)
            pel.append(r)
        if align is not None:
            p.alignment = align
        return p

    def add_blank():
        append_el(copy.deepcopy(blank_p))

    def add_image(path):
        full = path if os.path.isabs(path) else os.path.join(base_dir, path)
        if not os.path.exists(full):
            raise FileNotFoundError("图片不存在: %s" % full)
        p = add_para(body_ppr, body_rpr, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run().add_picture(full, width=Cm(img_width_cm))

    n_img = 0
    for i, blk in enumerate(blocks):
        if blk[0] == "heading":
            if i > 0:
                for _ in range(BLANKS_BETWEEN_SECTIONS):
                    add_blank()
            add_para(head_ppr, head_rpr, blk[1])
        elif blk[0] == "para":
            add_para(body_ppr, body_rpr, blk[1])
        elif blk[0] == "image":
            add_image(blk[1])
            add_para(body_ppr, body_rpr, blk[2], align=WD_ALIGN_PARAGRAPH.CENTER)
            n_img += 1
    for _ in range(BLANKS_BETWEEN_SECTIONS):
        add_blank()

    doc.save(out_path)
    return n_img, len(blocks)


def main():
    ap = argparse.ArgumentParser(
        description="把 md 写回 docx（沿用母本版式）",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    ap.add_argument("--md", default="CN-SCIENCE-CASE.md", help="输入的 markdown 文件")
    ap.add_argument("--template", default="CN-SCIENCE-CASE.docx",
                    help="提供版式的 docx 母本（读它取字体/字号/页边距）")
    ap.add_argument("--out", default="CN-SCIENCE-CASE.docx", help="输出的 docx 文件")
    ap.add_argument("--img-width-cm", type=float, default=10.0, help="图片宽度 (cm)")
    ap.add_argument("--backup-dir", default="backup", help="备份目录（不存在则创建）")
    ap.add_argument("--no-backup", action="store_true", help="覆盖前不备份")
    args = ap.parse_args()

    base_dir = os.path.dirname(os.path.abspath(args.md)) or "."
    if not os.path.exists(args.md):
        sys.exit("找不到 md: %s" % args.md)

    if os.path.exists(args.out) and not args.no_backup:
        os.makedirs(args.backup_dir, exist_ok=True)
        stamp = dt.datetime.now().strftime("%Y.%m.%d-%H%M")
        stem = os.path.splitext(os.path.basename(args.out))[0]
        bak = os.path.join(args.backup_dir, "%s.backup-%s.docx" % (stem, stamp))
        shutil.copy2(args.out, bak)
        print("已备份: %s" % bak)

    n_img, n_blk = build(args.template, args.md, args.out, args.img_width_cm, base_dir)
    print("已写出: %s（%d 个内容块，嵌入 %d 张图，图宽 %.1f cm）"
          % (args.out, n_blk, n_img, args.img_width_cm))


if __name__ == "__main__":
    main()
